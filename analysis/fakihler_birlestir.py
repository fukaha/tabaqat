"""Derin okuma çıktılarını (kaynak/analiz/okuma/parti_NN.json) Excel veri setine dönüştürür.

Girdi: kaynak/analiz/dosya/parti_NN.json (kişi dosyaları) ve kaynak/analiz/okuma/parti_NN.json (okuma kayıtları;
şema kaynak/analiz/YONERGE.md'de). Çıktı: kaynak/analiz/hanefi_fakihler_620_750.xlsx
    python analysis/fakihler_birlestir.py
"""
from __future__ import annotations

import glob
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
A = ROOT / "kaynak" / "analiz"
CELL_MAX = 32000   # Excel hücre sınırı 32.767


def CE(h: int) -> int:
    return int(622.54 + (h - 1) * 0.970224 + 0.485)


def flat(v):
    """Ajan çıktısındaki beklenmedik iç içe değerleri hücreye yazılabilir metne çevirir."""
    if isinstance(v, dict):
        return "; ".join(f"{k}: {flat(x)}" for k, x in v.items() if x not in (None, "", [], {}))
    if isinstance(v, (list, tuple)):
        return "; ".join(str(flat(x)) for x in v if x not in (None, ""))
    return v


def hm(h):
    if isinstance(h, dict):
        h = h.get("hicri")
    try:
        h = int(h)
    except (TypeError, ValueError):
        return None
    return f"{h}/{CE(h)}"


def j(xs, sep="; "):
    xs = [x for x in xs if x]
    return sep.join(xs) if xs else None


def g(d, *ks):
    for k in ks:
        d = (d or {}).get(k) if isinstance(d, dict) else None
    return d if d not in ("", [], {}) else None


def load():
    dos, oku = {}, {}
    for f in sorted(glob.glob(str(A / "dosya" / "parti_*.json"))):
        for d in json.load(open(f, encoding="utf-8")):
            dos[d["kimlik"]] = d
    for f in sorted(glob.glob(str(A / "okuma" / "parti_*.json"))):
        for o in json.load(open(f, encoding="utf-8")):
            oku[o["kimlik"]] = o
    return dos, oku


def sort_key(o, d):
    h = g(o, "vefat", "hicri")
    if h:
        return int(h)
    r = g(o, "vefat_donemi", "yaklasik_hicri") or ""
    m = re.findall(r"\d{3}", str(r))
    return sum(map(int, m)) / len(m) if m else (d.get("vefat_tahmini_sitede") or 999)


def rows(dos, oku):
    keep, out, L = [], [], {k: [] for k in ("yer", "hoca", "talebe", "eser", "gorev", "olay")}
    for pid, d in dos.items():
        o = oku.get(pid)
        if not o:
            continue
        if o.get("fakih") is False or o.get("kapsam_disi"):
            out.append({"kimlik": pid, "ad": o.get("tam_adi") or d["mevcut_turkce_ad"], "Arapça adı": d["arapca_ad"],
                        "vefat": hm(g(o, "vefat", "hicri")) or g(o, "vefat_donemi", "donem"),
                        "çıkarılma sebebi": "fakih değil" if o.get("fakih") is False else "kapsam dışı",
                        "gerekçe": o.get("fakih_gerekce") if o.get("fakih") is False else o.get("kapsam_gerekce"),
                        "kaynaklar": j(s["atif"] for s in d["kaynaklar"])})
            continue
        keep.append((sort_key(o, d), pid))
    keep.sort()
    R = []
    nmax = max((len(dos[p]["kaynaklar"]) for _, p in keep), default=0)
    for no, (_, pid) in enumerate(keep, 1):
        d, o = dos[pid], oku[pid]
        ad = o.get("meshur_adi") or o.get("tam_adi")
        vd = o.get("vefat_donemi") or {}
        eser = o.get("eserler") or []
        for x in o.get("bulundugu_yerler") or []:
            L["yer"].append({"kimlik": pid, "ad": ad, "yer": x.get("yer"), "tür": x.get("nitelik"), "kaynak": x.get("kaynak")})
        for x in o.get("seyahatler") or []:
            L["yer"].append({"kimlik": pid, "ad": ad, "yer": x.get("yer"), "tür": "seyahat" + (f" ({x['amac']})" if x.get("amac") else ""),
                             "kaynak": x.get("kaynak")})
        if o.get("dogum_yeri"):
            L["yer"].append({"kimlik": pid, "ad": ad, "yer": g(o, "dogum_yeri", "yer"), "tür": "doğum", "kaynak": g(o, "dogum_yeri", "kaynak")})
        for x in o.get("nisbe_yerleri") or []:
            L["yer"].append({"kimlik": pid, "ad": ad, "yer": x.get("yer"), "tür": f"nisbe ({x.get('nisbe') or ''})", "kaynak": x.get("kaynak")})
        for k, key in (("hoca", "hocalar"), ("talebe", "talebeler")):
            for x in o.get(key) or []:
                L[k].append({"kimlik": pid, "ad": ad, k: x.get("ad"), "ilim": x.get("ilim"), "kaynak": x.get("kaynak")})
        for x in eser:
            L["eser"].append({"kimlik": pid, "ad": ad, "vefat": hm(g(o, "vefat", "hicri")) or vd.get("donem"),
                              "eser": x.get("ad"), "Arapça adı": x.get("ad_ar"), "türü": x.get("tur"),
                              "ilim dalı": x.get("ilim_dali"), "dayandığı eser": x.get("dayandigi_eser"), "kaynak": x.get("kaynak")})
        for x in o.get("gorevler") or []:
            L["gorev"].append({"kimlik": pid, "ad": ad, "görev": x.get("gorev"), "yer": x.get("yer"), "kaynak": x.get("kaynak")})
        for x in o.get("siyasi_sosyal_olaylar") or []:
            L["olay"].append({"kimlik": pid, "ad": ad, "vefat": hm(g(o, "vefat", "hicri")) or vd.get("donem"),
                              "olay": x.get("olay"), "etkisi": x.get("etkisi"), "kanıt": x.get("kanit"), "kaynak": x.get("kaynak")})
        it = o.get("intikal") or {}
        r = {
            "no": no, "kimlik": pid,
            "tam adı": o.get("tam_adi"), "meşhur adı": o.get("meshur_adi"), "Arapça adı": d["arapca_ad"],
            "lakap ve unvanlar": j(o.get("lakap_ve_unvanlar") or [], ", "),
            "doğum (h./m.)": hm(g(o, "dogum", "hicri")), "doğum tarihi": g(o, "dogum", "tarih_metni"),
            "doğum kaynağı": g(o, "dogum", "kaynak"),
            "vefat (h./m.)": hm(g(o, "vefat", "hicri")), "vefat tarihi": g(o, "vefat", "tarih_metni"),
            "vefat kaynağı": g(o, "vefat", "kaynak"), "vefat ihtilafı": g(o, "vefat", "ihtilaf"),
            "vefat dönemi (tahmin)": vd.get("donem"), "tahminî aralık (h.)": vd.get("yaklasik_hicri"),
            "tahmin gerekçesi": vd.get("gerekce"), "tahmin güveni": vd.get("guven"),
            "nisbe yerleri": j(f"{x.get('yer')}" + (f" ({x['nisbe']})" if x.get("nisbe") else "") for x in o.get("nisbe_yerleri") or []),
            "doğum yeri": j([g(o, "dogum_yeri", "yer")] + ([f"({g(o, 'dogum_yeri', 'aciklama')})"] if g(o, "dogum_yeri", "aciklama") else []), " "),
            "yaşadığı bölge": j(o.get("yasadigi_bolge") or [], ", "),
            "bulunduğu şehir/köy": j(f"{x.get('yer')}" + (f" ({x['nitelik']})" if x.get("nitelik") else "") for x in o.get("bulundugu_yerler") or []),
            "seyahatleri": j(f"{x.get('yer')}" + (f" ({x['amac']})" if x.get("amac") else "") for x in o.get("seyahatler") or []),
            "intikal/göç": it.get("durum"), "göç güzergâhı": it.get("guzergah"), "göç sebebi": it.get("sebep"),
            "göç kanıtı": j([it.get("kanit"), f"[{it['kaynak']}]" if it.get("kaynak") else None], " "),
            "hocaları": j(f"{x.get('ad')}" + (f" ({x['ilim']})" if x.get("ilim") else "") for x in o.get("hocalar") or []),
            "talebeleri": j(f"{x.get('ad')}" + (f" ({x['ilim']})" if x.get("ilim") else "") for x in o.get("talebeler") or []),
            "eserleri": j(f"{x.get('ad')}" + (f" ({x['tur']})" if x.get("tur") else "") for x in eser),
            "eser türleri": j(sorted({x.get("tur") for x in eser if x.get("tur")}), ", "),
            "eserlerin ilim dalları": j(sorted({x.get("ilim_dali") for x in eser if x.get("ilim_dali")}), ", "),
            "görevleri": j(f"{x.get('gorev')}" + (f" ({x['yer']})" if x.get("yer") else "") for x in o.get("gorevler") or []),
            "ilmî ilgi alanları": j(o.get("ilmi_ilgi_alanlari") or [], ", "),
            "siyasî/sosyal olaylar": j(f"{x.get('olay')}" + (f": {x['etkisi']}" if x.get("etkisi") else "") for x in o.get("siyasi_sosyal_olaylar") or []),
            "olası tekrar kaydı": j(o.get("olasi_tekrar") or []),
            "fakihlik gerekçesi": flat(o.get("fakih_gerekce")),
            "notlar": o.get("notlar"),
            "kaynaklar": j(s["atif"] for s in d["kaynaklar"]) ,
            "sitede": f"https://fukaha.github.io/tabaqat/#/p/{pid}",
        }
        for i in range(nmax):
            s = d["kaynaklar"][i] if i < len(d["kaynaklar"]) else None
            txt = None
            if s:
                txt = f"{s['atif']}\n{s['madde_basligi']}\n\n{s['metin']}"
                if s.get("muhakkik_dipnotu"):
                    txt += f"\n\n— Muhakkik dipnotları —\n{s['muhakkik_dipnotu']}"
                if len(txt) > CELL_MAX:
                    txt = txt[:CELL_MAX] + "\n[… hücre sınırı nedeniyle kesildi]"
            r[f"tam metin {i + 1}"] = txt
        R.append(r)
    return R, out, L


GUIDE = [
    ("Kapsam", "Vefatı hicrî 620-750 (m. 1223-1349) arasında olan ve tabakāt eserlerinde maddesi bulunan Hanefî fakihler. "
               "Fıkıhla meşgul olduğu maddesinden anlaşılmayan kişiler (emîr, şair, tabib, yalnız muhaddis vb.) ve okuma sonunda "
               "vefatı bu aralığın dışına düşenler veri setinden çıkarılmış, kontrol için “Çıkarılanlar” sayfasında listelenmiştir."),
    ("Yöntem", "Her kişinin bütün tabakāt maddeleri tam metinden okunarak alanlar doldurulmuştur (yönerge: kaynak/analiz/YONERGE.md). "
               "Bilgi bulunmayan hücreler boş bırakılmıştır. Bilginin geçtiği kaynak uzun tablolarda her satırda verilmiştir."),
    ("Vefat dönemi", "Vefatı kaynaklarda kayıtlı olmayanlar için hoca ve talebelerinin vefatları, görev yaptığı hükümdarlar ve tarihli "
                     "olaylar esas alınarak dönem tahmini yapılmıştır: ilk çeyreği (1-25), ortaları (26-75), son çeyreği (76-90), "
                     "sonları (91-100). Gerekçe ve güven derecesi yanındaki sütunlardadır."),
    ("Atıf biçimi", "TDV İslâm Ansiklopedisi kaynakça üslûbu: müellif kısa adı, eser kısa adı, cilt (Roma rakamı), sayfa; tek ciltli "
                    "eserlerde “s.”. Madde numarası parantez içindedir."),
    ("Tam metinler", "Son sütunlarda kişinin her kaynaktaki maddesinin tam metni (madde metni + muhakkik dipnotları) kontrol için verilmiştir."),
    ("Telif", "Tam metinler telifli neşirlerdendir; dosya yalnız araştırma ekibi içindir, yayımlanmamalıdır."),
]


def write(R, out, L, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    head = PatternFill("solid", fgColor="1F3A5F")
    txtfill = PatternFill("solid", fgColor="F3EFE6")
    thin = Side(style="thin", color="D0D7E2")
    wb = Workbook()
    g_ = wb.active
    g_.title = "Kılavuz"
    g_["A1"] = "7./13. Yüzyıl Hanefî Fakihleri Veri Seti (h. 620-750)"
    g_["A1"].font = Font(name=F, size=14, bold=True, color="1F3A5F")
    for i, (k, v) in enumerate(GUIDE, 3):
        g_.cell(i, 1, k).font = Font(name=F, bold=True)
        g_.cell(i, 1).alignment = Alignment(vertical="top")
        c = g_.cell(i, 2, v)
        c.font = Font(name=F)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    g_.column_dimensions["A"].width = 18
    g_.column_dimensions["B"].width = 110
    r0 = len(GUIDE) + 5
    g_.cell(r0, 1, "Sayımlar").font = Font(name=F, bold=True, size=12, color="1F3A5F")
    keys = list(R[0].keys()) if R else []
    n = len(R) + 1
    counts = [("fakih (veri setinde)", f'=COUNTA(Fakihler!B2:B{n})'),
              ("çıkarılan", f'=COUNTA(Çıkarılanlar!A2:A{len(out) + 1})')]
    for col in ("doğum (h./m.)", "vefat (h./m.)", "vefat dönemi (tahmin)", "doğum yeri", "yaşadığı bölge", "bulunduğu şehir/köy",
                "intikal/göç", "hocaları", "eserleri", "görevleri", "siyasî/sosyal olaylar"):
        if col in keys:
            Lc = get_column_letter(keys.index(col) + 1)
            counts.append((f"dolu: {col}", f'=COUNTIF(Fakihler!{Lc}2:{Lc}{n},"?*")'))
    if "intikal/göç" in keys:
        Lc = get_column_letter(keys.index("intikal/göç") + 1)
        counts.append(("intikal: evet", f'=COUNTIF(Fakihler!{Lc}2:{Lc}{n},"evet")'))
    for i, (k, f) in enumerate(counts, r0 + 1):
        g_.cell(i, 1, k).font = Font(name=F)
        g_.cell(i, 2, f).font = Font(name=F)
        g_.cell(i, 2).alignment = Alignment(horizontal="left")

    WIDE = {"tam adı": 36, "meşhur adı": 22, "Arapça adı": 36, "lakap ve unvanlar": 22, "hocaları": 40, "talebeleri": 34,
            "eserleri": 44, "kaynaklar": 50, "notlar": 44, "fakihlik gerekçesi": 34, "tahmin gerekçesi": 40,
            "bulunduğu şehir/köy": 30, "seyahatleri": 26, "görevleri": 34, "siyasî/sosyal olaylar": 40, "göç kanıtı": 36,
            "vefat ihtilafı": 30, "olası tekrar kaydı": 14, "kanıt": 50, "etkisi": 40, "gerekçe": 50, "kaynak": 40, "eser": 34, "doğum kaynağı": 30,
            "vefat kaynağı": 30, "göç güzergâhı": 26, "nisbe yerleri": 24, "doğum yeri": 20, "sitede": 16}
    AR = {"Arapça adı", "kanıt", "göç kanıtı"}

    def sheet(title, data, freeze="C2"):
        ws = wb.create_sheet(title)
        if not data:
            return
        ks = list(data[0].keys())
        for c, k in enumerate(ks, 1):
            cell = ws.cell(1, c, k)
            cell.font = Font(name=F, bold=True, color="FFFFFF")
            cell.fill = head
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = 60 if k.startswith("tam metin") else WIDE.get(k, 13)
        for r, d in enumerate(data, 2):
            for c, k in enumerate(ks, 1):
                v = d.get(k)
                v = flat(v)
                cell = ws.cell(r, c, v if v not in ("", None) else None)
                cell.font = Font(name=F, size=10)
                cell.border = Border(bottom=thin)
                full = k.startswith("tam metin")
                cell.alignment = Alignment(wrap_text=True, vertical="top",
                                           horizontal="right" if (k in AR or full) else None,
                                           readingOrder=2 if (k in AR or full) else 0)
                if full and v:
                    cell.fill = txtfill
                if k == "sitede" and v:
                    cell.hyperlink = v
                    cell.font = Font(name=F, size=10, color="1F5FBF", underline="single")
            if any(k.startswith("tam metin") for k in ks):
                ws.row_dimensions[r].height = 120
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32

    sheet("Fakihler", R, "D2")
    sheet("Yerler", L["yer"])
    sheet("Hocalar", L["hoca"])
    sheet("Talebeler", L["talebe"])
    sheet("Eserler", L["eser"])
    sheet("Görevler", L["gorev"])
    sheet("Siyasî ve sosyal olaylar", L["olay"])
    sheet("Çıkarılanlar", out)
    wb.save(path)


if __name__ == "__main__":
    dos, oku = load()
    R, out, L = rows(dos, oku)
    write(R, out, L, A / "hanefi_fakihler_620_750.xlsx")
    eksik = [k for k in dos if k not in oku]
    print(f"okunan {len(oku)}/{len(dos)}, veri setinde {len(R)}, çıkarılan {len(out)}, okunmamış {len(eksik)}")
